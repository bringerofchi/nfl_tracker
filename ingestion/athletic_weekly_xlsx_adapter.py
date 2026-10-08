"""
The Athletic weekly projections xlsx adapter.

New specimen (Week 2, 2026): unlike the season-long Athletic workbook this
adapter's sibling (athletic_xlsx_adapter.py) was built against -- one sheet
per position, "Player"/"TM"/raw-stat-abbreviation columns -- this file is a
single sheet with QB/RB/WR/TE laid out as four side-by-side column blocks,
named like "NFL_2026_Week_<N>_Half_PPR_Weekly". Per the frozen boundary this
file is responsible ONLY for reading the workbook and producing
NormalizedObservations (scope='weekly'); everything else (identity
resolution, validation, confidence routing, versioning) is the unchanged
ingestion.pipeline.ingest_collector_batch() -> _route_field() path every
other source goes through.

Known limitations, stated plainly:
- Raw counting stats only. The sheet's own "FPS" column (its precomputed
  half-PPR point total) is deliberately NOT imported -- same principle as
  the season adapter: this app always derives fantasy points itself, never
  trusts a source's own total. The sheet name says "Half_PPR" but that
  label only describes how FPS was computed; since FPS itself is skipped,
  nothing half-PPR-specific ever reaches the trusted projections table.
- fumbles_lost is not present in this workbook for any position; that
  observation is simply never produced, same as the season adapter.
- The sheet name (and therefore the week number) is not parsed to confirm
  it matches the requested week_number -- this adapter trusts the caller,
  same as every other adapter's fetch(season_id, week_number) contract.
"""
from typing import List

from ingestion.collector import SourceAdapter, NormalizedObservation, AdapterUnavailable

# Block header label (row 1) -> stat map for that position's block (row 2 = field names).
_QB_STAT_MAP = {"Pass YD": "pass_yards", "Pass TD": "pass_tds", "INT": "interceptions",
                "Rush Att": "rush_attempts", "Rush YD": "rush_yards", "Rush TD": "rush_tds"}
_RB_STAT_MAP = {"Rush Att": "rush_attempts", "Rush YD": "rush_yards", "Rush TD": "rush_tds",
                "REC": "receptions", "REC YD": "receiving_yards", "REC TD": "receiving_tds"}
_WR_STAT_MAP = {"Rush YD": "rush_yards", "Rush TD": "rush_tds",
                "REC": "receptions", "REC YD": "receiving_yards", "REC TD": "receiving_tds"}
_TE_STAT_MAP = {"REC": "receptions", "REC YD": "receiving_yards", "REC TD": "receiving_tds"}

_POSITION_STAT_MAPS = {"QB": _QB_STAT_MAP, "RB": _RB_STAT_MAP, "WR": _WR_STAT_MAP, "TE": _TE_STAT_MAP}

# Same rationale as the season adapter: real structured spreadsheet cells,
# not OCR/vision uncertainty -- but still a personal analyst's projection,
# and this adapter has no per-field signal of its own.
_PROJECTION_CONFIDENCE = 0.9


class AthleticWeeklyXlsxAdapter(SourceAdapter):
    """Reads the single-sheet, four-position-block weekly workbook and emits
    scope='weekly' raw-stat projections for QB/RB/WR/TE. fetch()'s
    season_id argument is accepted for interface compatibility but not
    otherwise used (the workbook itself is already scoped to one week)."""

    def __init__(self, source_name: str, file_path: str):
        self.source_name = source_name
        self.file_path = file_path

    def fetch(self, season_id: int, week_number: int) -> List[NormalizedObservation]:
        try:
            import openpyxl
        except ImportError as e:
            raise AdapterUnavailable("openpyxl is not installed (pip install openpyxl)", transient=False) from e

        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
        except FileNotFoundError as e:
            raise AdapterUnavailable(f"Source file not found: {self.file_path}", transient=True) from e
        except Exception as e:
            raise AdapterUnavailable(f"Could not open workbook: {e}", transient=False) from e

        if not wb.sheetnames:
            raise AdapterUnavailable("Workbook has no sheets", transient=False)
        ws = wb[wb.sheetnames[0]]

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise AdapterUnavailable("Workbook has no header rows", transient=False)
        position_row, field_row = rows[0], rows[1]

        # Find each position block's starting column (first occurrence only --
        # same "first group, skip duplicates" discipline as the season adapter).
        block_starts = {}
        for i, label in enumerate(position_row):
            if label in _POSITION_STAT_MAPS and label not in block_starts:
                block_starts[label] = i

        # Every labelled section start in row 1, including sections this adapter
        # ignores (Flex / Superflex / DST appear from Week 5 on). A block must end
        # at the NEXT section of any kind, otherwise the last known block (TE)
        # runs into the ignored sections and their duplicate "Name"/"Team"
        # headers overwrite TE's own columns (Week 5 specimen read DST names as TEs).
        all_section_starts = sorted(i for i, label in enumerate(position_row) if label)

        observations: List[NormalizedObservation] = []
        block_items = sorted(block_starts.items(), key=lambda kv: kv[1])
        for idx, (position, start_col) in enumerate(block_items):
            later_starts = [c for c in all_section_starts if c > start_col]
            end_col = later_starts[0] if later_starts else len(field_row)
            block_fields = list(field_row[start_col:end_col])
            col_index = {name: start_col + j for j, name in enumerate(block_fields) if name is not None}
            if "Name" not in col_index or "Team" not in col_index:
                continue  # block doesn't have the expected shape -- skip rather than guess
            stat_map = _POSITION_STAT_MAPS[position]

            for row in rows[2:]:
                player_name = row[col_index["Name"]]
                if not player_name:
                    continue
                team_hint = row[col_index["Team"]] if col_index["Team"] < len(row) else None
                for header_key, stat_name in stat_map.items():
                    if header_key not in col_index or col_index[header_key] >= len(row):
                        continue
                    value = row[col_index[header_key]]
                    if value is None:
                        continue
                    observations.append(NormalizedObservation(
                        player_name_raw=str(player_name),
                        position_hint=position,
                        data_type="projection",
                        stat_name=stat_name,
                        value=round(float(value), 2),
                        confidence=_PROJECTION_CONFIDENCE,
                        week_number=week_number,
                        scope="weekly",
                        team_hint=team_hint,
                    ))
        return observations
